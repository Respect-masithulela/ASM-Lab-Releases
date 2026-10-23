"""the source tool input exports -> ASM Lab project (.asmlab.json). Generic version of sample_test_1/source_to_asm.py.

Usage:  py network_export_to_asmlab.py <folder> <file prefix, e.g. "Network1"> "<project name>"
Needs in <folder>: "<prefix> branches.xlsx", "<prefix> nodes.xlsx", "<prefix> tunnels.xlsx"; optional "<prefix> shafts / duct /
control managers (or cm) / fans / dev end.xlsx", and "fan_curves.json" for the fans: {"Main fan": [[Q, P], ...], "11kW": [[Q, P], ...]}
(the source tool does not export library fan curves: copy the points from the source tool Libraries > Fans).

Device mapping (aero only; heat devices are carried as negligible resistance and listed in the notes):
  Tunnel / Shaft / Duct      geometry, Atkinson k, the source tool length (ducts are ordinary circular branches in the source tool and here)
  Development heading        dead-end airway (the source tool's internal heading ventilation is not carried over)
  User-defined resistance    resistanceOverride R
  Tunnel reduced area        orifice R = rho / (2 (0.65 A_open)^2)   [assumption: the source tool's law is not exported]
  Door "P Pa @ Q m3/s"       R = P / Q^2
  Fixed flow                 fixedQ + fixedFlow regulator
  Fixed press. difference    constant pressure loss P in the airway's direction (fan with a flat curve of -P)
  User-defined heatload "@ P Pa"   as fixed pressure difference P
  Cooler / refrigeration / hoist / pump / heat   R = 1e-4 (heat only)
  Fan                         curve from fan_curves.json (fan library name); if missing, the fan is written with no curve and flagged
Boundaries: dead-end nodes at the surface (depth <= 5 m) = surface connections at the source tool's barometric pressure; the intake/exhaust role
comes from the branch's IntakeReturn field.
"""
import json, math, re, os, sys, collections
import openpyxl

folder, pre, pname = sys.argv[1], sys.argv[2], sys.argv[3]
# optional key=value options (benchmark exports, exports_by_category/<cat>/<example>/the source tool/{Input,Output}):
#   results=<the source tool Output folder>  control managers out.xlsx (R of devices; Q of fixed pressure differences), nodes out.xlsx
#   temps=results                  node temperatures from the source tool's solved nodes out (aero category: density as the source tool used it)
#   curves=<fan curve json>        instead of <folder>/fan_curves.json; names are matched ignoring case, spaces and symbols
#   out=<folder>                   where the project is written;  noduct=1  leaves ducts to ducts_to_asmlab.py
opt = dict(a.split("=", 1) for a in sys.argv[4:])
# tier=blind: the source tool INPUTS ONLY. No device R, fixed-pressure flows or temperatures from the source tool's solution; devices from their own definitions
# (door P @ Q, measured P @ Q, orifice law for reduced areas and ducts through walls, constant pressure loss for fixed pressure
# differences and coolers with a stated pressure drop), kg/s fixed flows at the inlet node's input density. Production zones keep
# the source tool's solved R (their inputs are not in the exported tables).
BLIND = opt.get("tier") == "blind"
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from psychro import rho as _rho
outdir = opt.get("out", folder)


def L(*names):
    for n in names:
        for p in (os.path.join(folder, f"{pre} {n}.xlsx"), os.path.join(folder, f"{n}.xlsx")):
            if os.path.exists(p): return list(openpyxl.load_workbook(p, read_only=True, data_only=True).active.iter_rows(values_only=True))[1:]
    return []


num = lambda s: float(re.findall(r"-?[\d.]+", str(s))[0])


def geom(g):
    m = re.match(r"\(Irregular\)\s*([\d.]+) m \* ([\d.]+) m", str(g))
    if m: return "irregular", float(m[1]), float(m[2])  # perimeter, area
    m = re.match(r"\(Rect\)\s*([\d.]+) m X ([\d.]+) m", str(g))
    if m: return "rectangular", float(m[2]), float(m[1])  # the source tool writes "(Rect) H m X W m": height first
    m = re.match(r"\(Circ\)\s*([\d.]+) m", str(g))
    return "circular", float(m[1]), float(m[1])


br, nodes = L("branches", "input"), L("nodes")
# a few rows of the source tool's grid copy (Ctrl+C, used where its Excel export crashes) lose cells: realign them
nodes = [r if isinstance(r[0], (int, float)) else (None,) + tuple(r) for r in nodes]
_nn = {r[1] for r in nodes}
for i, r in enumerate(br):
    if r[5] not in _nn or r[6] not in _nn:
        j = next((j for j in range(2, len(r) - 1) if r[j] in _nn and r[j + 1] in _nn), None)
        if j is not None: br[i] = tuple(r[:j]) + (None,) * (5 - j) + tuple(r[j:]) if j < 5 else tuple(r[:5]) + tuple(r[j:])
_tun = L("tunnels", "tunnel")
tun = {r[0]: r for r in _tun if isinstance(r[0], (int, float))}
_lost = [(None,) + tuple(r) for r in _tun if not isinstance(r[0], (int, float))]  # rows that lost their ID: match by name
for r in br:
    if r[2] == "Tunnel" and r[0] not in tun:
        m = [t for t in _lost if t[1] == r[1]]
        if m: tun[r[0]] = m[0]; _lost.remove(m[0])
sh = {r[0]: r for r in L("shafts")}; du = {r[0]: r for r in L("duct", "ducts")}
cm = {r[0]: r for r in L("control managers", "cm")}; fans = {r[0]: r for r in L("fans")}; dev = {r[0]: r for r in L("dev end", "development headings", "dev headings")}
_cv = opt.get("curves", os.path.join(folder, "fan_curves.json"))
curves = json.load(open(_cv, encoding="utf-8")) if os.path.exists(_cv) else {}
_key = lambda s: re.sub(r"[^0-9a-z.]", "", str(s).lower())  # keep the decimal point: "7.5kW" is not "75kW"
_ck = {_key(k): k for k in curves if not k.startswith("_")}


def curve(f):
    """library curve of a the source tool fan row, scaled for speed ratio (Q ~ n, P ~ n^2) and cluster (parallel: Q x N, series: P x N)"""
    lib = _ck.get(_key(f[2]))
    if lib is None: return None, f[2], 1.2
    s = float(f[4] or 100) / 100; n = int(f[5] or 1); par = str(f[6]).lower().startswith("par"); ser = str(f[6]).lower().startswith("ser")
    pts = [[q * s * (n if par else 1), p * s * s * (n if ser else 1)] for q, p in curves[lib]]
    if s != 1 or par or ser: notes[f"fan '{lib}' scaled: speed {s:g}, {n} in {f[6]}"] += 1
    return pts, lib, curves.get("_rhoRef", {}).get(lib, 1.2)


nd = {r[1]: r for r in nodes}
_cmp = os.path.join(opt["results"], "control managers out.xlsx") if "results" in opt else os.path.join(folder, "CM results.xlsx")
cmres = {r[0]: (float(r[4]), abs(float(r[7])), float(r[8] or 0) if len(r) > 8 and isinstance(r[8], (int, float)) else 0.0) for r in list(openpyxl.load_workbook(_cmp, read_only=True, data_only=True).active.iter_rows(values_only=True))[1:] if r[0] is not None and r[4] is not None} if os.path.exists(_cmp) else {}
dpq = json.load(open(os.path.join(folder, "fixed_dp_flows.json"))) if os.path.exists(os.path.join(folder, "fixed_dp_flows.json")) else {}
if "results" in opt and not dpq and not BLIND:  # the source tool flow through each fixed pressure difference, from the solved control managers
    dpq = {str(b): abs(q) for b, (q, _, _) in cmres.items() if b in cm and ("Fixed press" in str(cm[b][2]) or "heatload" in str(cm[b][2])) and abs(q) > 1e-3}
# production zones (stopes) have no control-manager row: their R from the source tool's branch results (ID, Name, Type, Q, M, v, dP, R)
_bo = os.path.join(opt.get("results", ""), "branches out.xlsx")
if "results" in opt and os.path.exists(_bo):
    for r in list(openpyxl.load_workbook(_bo, read_only=True, data_only=True).active.iter_rows(values_only=True))[1:]:
        if r[0] not in cmres and str(r[2]) == "Production zone" and isinstance(r[3], (int, float)) and isinstance(r[7], (int, float)):
            cmres[r[0]] = (float(r[3]), abs(float(r[6] or 0)), float(r[7]))
_fo = os.path.join(opt.get("results", ""), "fans out.xlsx")
_fanout = {r[0]: (float(r[4]), float(r[6])) for r in list(openpyxl.load_workbook(_fo, read_only=True, data_only=True).active.iter_rows(values_only=True))[1:]
           if isinstance(r[4], (int, float)) and isinstance(r[6], (int, float))} if "results" in opt and os.path.exists(_fo) else {}
tres = {}
if opt.get("temps") == "results":  # solved node conditions (the source tool nodes out: Name, Pressure, Wet bulb, Dry bulb, ...)
    _rows = list(openpyxl.load_workbook(os.path.join(opt["results"], "nodes out.xlsx"), read_only=True, data_only=True).active.iter_rows(values_only=True))
    _h = [str(c).split("\n")[0].strip().lower() for c in _rows[0]]
    _n, _w, _d = _h.index("name"), _h.index("wet bulb"), _h.index("dry bulb")
    tres = {r[_n]: (float(r[_w]), float(r[_d])) for r in _rows[1:] if r[_n] is not None and isinstance(r[_d], (int, float))}
deg = collections.Counter()
for r in br: deg.update([r[5], r[6]])
role = {}
aw, regs, fl, notes, cnt = [], [], [], collections.Counter(), collections.Counter()
# surface = within 5 m of the shallowest node (the source tool depths can be measured from a datum below the collar, e.g. Metal mine at 216 m)
surf = max(5.0, min(float(nd[n][4]) for n in deg) + 5)
# a dead end on a control manager is also an atmosphere connection at any depth (the source tool flows it: Coal's two exhaust portals at 80 m,
# Metal's leakage-out at 2 062 m, Platinum node 366 at 840 m); dead-end tunnels and development headings below surface stay blind
for r in br:
    for n in (r[5], r[6]):
        if deg[n] == 1 and (float(nd[n][4]) <= surf or str(r[2]).startswith("Control manager")):
            role[n] = "intake" if str(r[4]).lower().startswith("intake") else "exhaust"
            if float(nd[n][4]) > surf: notes[f"deep control-manager dead end {n} ({r[2]}, {float(nd[n][4]):g} m) = atmosphere"] += 1

if not role and "StartNode" in nd: role["StartNode"] = "intake"; notes["closed surface loop: StartNode = atmosphere (fixed pressure)"] += 1
if not role:  # closed loop without StartNode: the 0 Pa fixed pressure difference closing the loop at surface; its inlet node is the reference
    zero = [r for r in br if r[0] in cm and "Fixed press" in str(r[2]) and num(cm[r[0]][5]) == 0]
    if zero:
        r = min(zero, key=lambda r: float(nd[r[5]][4])); role[r[5]] = "intake"
        notes[f"closed surface loop: inlet of 0 Pa link {r[0]} ({r[5]}) = atmosphere reference"] += 1
# one atmosphere: the source tool solves every surface connection at the same barometric pressure (the most common input value among them),
# elevation-corrected; a stale per-node input (Joel node 528: 86.0 against 85.548 kPa) would otherwise push air between boundaries
deep_bc = {n for n in role if n != "StartNode" and float(nd[n][4]) > surf}  # at their own the source tool barometric input (hydrostatic at 1.2 kg/m3 is kPa off at depth)
_bp = collections.Counter(round(float(nd[n][9]), 3) for n in role if n != "StartNode" and n not in deep_bc)
_z0 = {n: -float(nd[n][4]) for n in role if n not in deep_bc}
P_ATM = _bp.most_common(1)[0][0] * 1000 if _bp else None
Z_ATM = sum(_z0.values()) / len(_z0) if _z0 else 0.0
if len(_bp) > 1: notes[f"surface boundaries at one pressure {P_ATM:g} Pa (inputs differed: {sorted(_bp)})"] += 1
# matched tier: a boundary is held at the pressure the source tool's own solution holds it at (Joel 528 etc.: 86.000 kPa in the source tool's results, against the
# 85.548 kPa a single elevation-corrected atmosphere gives); the blind tier keeps the input / one-atmosphere rule
PV = {}
if "results" in opt and not BLIND and os.path.exists(os.path.join(opt["results"], "nodes out.xlsx")):
    _rows = list(openpyxl.load_workbook(os.path.join(opt["results"], "nodes out.xlsx"), read_only=True, data_only=True).active.iter_rows(values_only=True))
    _h = [str(c).splitlines()[0].strip().lower() for c in _rows[0]]
    PV = {r[_h.index("name")]: float(r[_h.index("pressure")]) * 1000 for r in _rows[1:] if r[_h.index("name")] is not None and isinstance(r[_h.index("pressure")], (int, float))}
out_nodes = []
for name in {n for r in br for n in (r[5], r[6])}:
    r = nd[name]
    wb, db = tres.get(name, (float(r[7]), float(r[8])))
    o = dict(id=name, x=-float(r[2]), y=float(r[3]), elevation=-float(r[4]), level=r[5], temperature=db, wetBulbTemp=wb,
             inputBaroKPa=float(r[9]), inputDB=float(r[8]), inputWB=float(r[7]))  # the source tool input air at the node (for MFIRE's single density)
    if name in deep_bc: o.update(type=role[name], fixedPressure=round(float(r[9]) * 1000, 1))
    elif name in role: o.update(type=role[name], fixedPressure=round((P_ATM or float(r[9]) * 1000) - 1.2 * 9.81 * (-float(r[4]) - Z_ATM), 1))
    if name in role and name in PV: o["fixedPressure"] = round(PV[name], 1)
    out_nodes.append(o)

for r in br:
    bid, name, typ, a, b = r[0], r[1], r[2], r[5], r[6]
    pa, pb = nd[a], nd[b]
    Lg = max(math.dist((pa[2], pa[3], -pa[4]), (pb[2], pb[3], -pb[4])), 0.5)
    o = dict(id=str(bid), name=f"{name} [{typ}]", fromNode=a, toNode=b, frictionFactor=0.012, shockLoss=0, airDensity=1.2, length=round(Lg, 3),
             shape="rectangular", width=5.0, height=5.0)
    det = cm[bid][5] if bid in cm else ""
    cnt[typ] += 1
    cr = cmres.get(bid)
    if (not BLIND or typ == "Production zone") and (typ.startswith("Control manager") or typ == "Production zone") and "Fixed flow" not in typ and "Specified flow" not in str(det) and "(Leakage) Between" not in typ and "(Leakage) Out" not in typ and not (typ == "Control manager" and "/s" in str(det)) and cr and abs(cr[0]) > 0.05:
        # the source tool's own resistance column where it has one (its pressure drop is rounded to whole Pa, which spoils dP/Q^2 at low flow)
        o["resistanceOverride"] = round(cr[2] if cr[2] > 0 else max(cr[1], 0.0) / cr[0] ** 2, 8) or 1e-6
        notes["control manager R = dP / Q^2 from the source tool results (CM results.xlsx)"] += 1
    elif bid in du:
        t = du[bid]; s, w, h = geom(t[3]); o.update(shape=s, width=w, height=h, diameter=w, frictionFactor=num(t[4]), length=float(t[2]), name=f"{name} [Duct]")
    elif typ == "Tunnel" and bid in tun:
        t = tun[bid]; s, w, h = geom(t[3]); o.update(shape=s, width=w, height=h, frictionFactor=num(t[4]), length=float(t[2]))
        if s == "circular": o["diameter"] = w
        if s == "irregular":  # the source tool perimeter and area: square of the same area, Atkinson R from the real perimeter
            o.update(shape="rectangular", width=round(math.sqrt(h), 3), height=round(math.sqrt(h), 3), resistanceOverride=round(o["frictionFactor"] * o["length"] * w / h ** 3, 8))
            notes["irregular tunnel: square of the source tool area, R = k L P / A^3 with the source tool perimeter"] += 1
    elif typ == "Shaft":
        t = sh[bid]; s, w, h = geom(t[3]); o.update(shape=s, width=w, height=h, diameter=w, frictionFactor=num(t[4]), length=float(t[2]))
        if s == "irregular":
            o.update(shape="rectangular", width=round(math.sqrt(h), 3), height=round(math.sqrt(h), 3), resistanceOverride=round(o["frictionFactor"] * o["length"] * w / h ** 3, 8)); o.pop("diameter")
            notes["irregular shaft: square of the source tool area, R = k L P / A^3 with the source tool perimeter"] += 1
    elif typ == "Development heading":
        if bid in dev: s, w, h = geom(dev[bid][4]); o.update(shape=s, width=w, height=h, length=float(dev[bid][3]))
        notes["development heading: dead end, the source tool heading ventilation not carried"] += 1
    elif typ == "Fan" and str(fans[bid][7]).lower() != "on":
        fq = _fanout.get(bid)  # a stopped fan: airway with the resistance the source tool gives it (its solved dP / Q^2)
        o["resistanceOverride"] = round(abs(fq[1]) / fq[0] ** 2, 6) if fq and abs(fq[0]) > 0.05 else 1.0
        notes["fan switched off -> airway with the source tool's off resistance (dP / Q^2)"] += 1
    elif typ == "Fan" and opt.get("fans") == "fixed" and _fanout.get(bid) and _fanout[bid][0] > 0.05:
        # fixed-flow test: every running fan is a fixed flow at the volume the source tool's own solution gives it (the source tool's solution is unchanged by this)
        q = round(_fanout[bid][0], 4); o["fixedQ"] = q
        regs.append(dict(id=f"R{bid}", airwayId=str(bid), name=f"Fan {name} as fixed flow {q:g} m3/s", mode="fixedFlow", status="partial", fixedFlow=q))
        notes["fan as fixed flow at the source tool's fan volume (fixed-flow test)"] += 1
    elif typ == "Fan":
        f = fans[bid]; o["resistanceOverride"] = 1e-6; o["device"] = "fan"
        pts, lib, rho = curve(f)
        fl.append(dict(id=f"F{bid}", airwayId=str(bid), name=f"{name} ({lib})", position=0.5, fanKind=curves.get("_kind", {}).get(lib, "main" if "main" in str(lib).lower() else "auxiliary"),
                       source="points", points=[dict(Q=q, P=p) for q, p in (pts or [])], speedRef=1, operatingSpeedRps=1 if str(f[7]).lower() == "on" else 0, rhoRef=rho))
        if not pts: notes[f"fan '{lib}': NO CURVE in the fan curve file"] += 1
    elif "(Leakage)Fixed resistance" in typ:
        o["resistanceOverride"] = max(num(det), 1e-6)
    elif ("(Leakage)" in typ or typ == "Control manager") and "/s" in str(det):  # untyped control manager with a flow = the source tool specified flow (Platinum 649: 18 m3/s against a 1 722 Pa rise)
        q = num(det)
        if "kg/s" in str(det):
            q = q / _rho(float(nd[a][9]), float(nd[a][8]), float(nd[a][7])) if BLIND else (abs(cr[0]) if cr and abs(cr[0]) > 1e-6 else q / 1.2)
            notes["leakage mass flow (kg/s) -> m3/s" + (" at input node density" if BLIND else " as the source tool's volume flow")] += 1
        o["fixedQ"] = q
        regs.append(dict(id=f"R{bid}", airwayId=str(bid), name=f"Leakage {q:g} m3/s ({typ.split(')')[-1].strip()})", mode="fixedFlow", status="partial", fixedFlow=q))
        notes[f"{typ}: specified flow as fixed flow"] += 1
    elif typ == "Production zone":
        o["resistanceOverride"] = 0.5; notes["production zone (stope): R = 0.5 placeholder until the source tool results give R = dP/Q^2"] += 1
    elif "User-defined resistance" in typ:
        o["resistanceOverride"] = max(num(det), 1e-6)
    elif "reduced area" in typ:
        s, w, h = geom(det.split(";")[0]); A = (math.pi * w * w / 4 if s == "circular" else w * h) * num(det.split("=")[1]) / 100
        o["resistanceOverride"] = round(1.2 / (2 * (0.65 * A) ** 2), 4)
    elif typ.endswith("Door"):
        P, Q = re.findall(r"[\d.]+", det)[:2]; o["resistanceOverride"] = round(float(P) / float(Q) ** 2, 4)
    elif "Fixed flow" in typ:
        # the source tool fixed flows are in m3/s or kg/s; a mass flow becomes the volume flow the source tool ran it at (M / rho of its own solution)
        q = num(det)
        if "kg/s" in str(det):
            q = q / _rho(float(nd[a][9]), float(nd[a][8]), float(nd[a][7])) if BLIND else (abs(cr[0]) if cr and abs(cr[0]) > 1e-6 else q / 1.2)
            notes["fixed mass flow (kg/s) -> m3/s" + (" at input node density" if BLIND else " as the source tool's volume flow")] += 1
        o["fixedQ"] = q
        regs.append(dict(id=f"R{bid}", airwayId=str(bid), name=f"Fixed flow {q:g} m3/s ({name})", mode="fixedFlow", status="partial", fixedFlow=q))
    elif ("Fixed press" in typ or "heatload" in typ) and str(bid) in dpq:
        P = num(det.split("@")[1]) if "@" in det else num(det)
        o["resistanceOverride"] = round(P / dpq[str(bid)] ** 2, 8)
        notes["fixed pressure difference as R = P / Q_the source tool^2"] += 1
    elif ("Fixed press" in typ or "heatload" in typ) and num(det.split("@")[1] if "@" in det else det) == 0:
        o["resistanceOverride"] = 1e-6; notes["0 Pa fixed pressure difference = open connection, R = 1e-6"] += 1
    elif "Fixed press" in typ or "heatload" in typ:
        P = num(det.split("@")[1]) if "@" in det else num(det)
        o["resistanceOverride"] = 1e-6; o["device"] = "fan"
        fl.append(dict(id=f"P{bid}", airwayId=str(bid), name=f"Fixed pressure difference {P:g} Pa ({name})", position=0.5, fanKind="booster", source="points",
                       points=[dict(Q=-1000, P=-P), dict(Q=1000, P=-P)], speedRef=1, operatingSpeedRps=1, rhoRef=1.2))
        notes["fixed pressure difference as a constant pressure loss in the airway direction"] += 1
    elif "Meas. flow" in typ and "@" in det:
        P, Q = re.findall(r"[\d.]+", det)[:2]; o["resistanceOverride"] = round(float(P) / float(Q) ** 2, 6); notes["measured P @ Q -> R = P / Q^2"] += 1
    elif "Duct through wall" in typ:
        d = num(det.split("Inlet")[1]) if "Inlet" in det else geom(det.split(";")[0])[1]; Ao = math.pi * d * d / 4
        o["resistanceOverride"] = round(1.5 * 1.2 / (2 * Ao ** 2), 6); notes["duct through wall: 1.5 velocity heads at the inlet diameter"] += 1
    elif typ.startswith("Control manager") and re.search(r"(?:@|Pressure drop:)\s*([\d.]+)\s*Pa", det) and float(re.search(r"(?:@|Pressure drop:)\s*([\d.]+)\s*Pa", det)[1]) >= 5:
        P = float(re.search(r"(?:@|Pressure drop:)\s*([\d.]+)\s*Pa", det)[1])
        o["resistanceOverride"] = 1e-6; o["device"] = "fan"
        fl.append(dict(id=f"P{bid}", airwayId=str(bid), name=f"{typ.split('| ')[-1]} pressure drop {P:g} Pa ({name})", position=0.5, fanKind="booster", source="points",
                       points=[dict(Q=-1000, P=-P), dict(Q=1000, P=-P)], speedRef=1, operatingSpeedRps=1, rhoRef=1.2))
        notes["device with a stated pressure drop -> constant pressure loss"] += 1
    elif typ.startswith("Control manager"):
        o["resistanceOverride"] = 1e-4; notes[f"{typ.split('| ')[-1]}: heat only, R = 1e-4"] += 1
    aw.append(o)

# the source tool's StartNode connections are virtual (not drawn in the source tool): replace near-zero-R links to StartNode by atmosphere boundaries
if role.get("StartNode"):
    P0 = round(float(nd["StartNode"][9]) * 1000, 1); links = [o for o in aw if "StartNode" in (o["fromNode"], o["toNode"]) and o.get("resistanceOverride", 1) <= 1e-5]
    if links and len(links) == sum("StartNode" in (o["fromNode"], o["toNode"]) for o in aw):
        aw = [o for o in aw if o not in links]; bn = {}
        for o in links: bn[o["toNode"] if o["fromNode"] == "StartNode" else o["fromNode"]] = "intake" if o["fromNode"] == "StartNode" else "exhaust"
        out_nodes = [n for n in out_nodes if n["id"] != "StartNode"]
        for n in out_nodes:
            if n["id"] in bn: n.update(type=bn[n["id"]], fixedPressure=P0)
        notes[f"StartNode connections ({len(links)}, R ~ 0) replaced by atmosphere boundaries at {P0:g} Pa"] += 1

# ---- ducts: the source tool duct branches (+ the fan in the chain) -> one ASM duct routed along the tunnels it serves
ducts_out = []
duct_ids = {str(b) for b in du}
if duct_ids and not opt.get("noduct"):
    import heapq
    A = {o["id"]: o for o in aw}
    dnode = collections.Counter()
    for i in duct_ids: dnode.update([A[i]["fromNode"], A[i]["toNode"]])
    fan_br = [o for o in aw if o.get("device") == "fan" and o["id"] in {str(b) for b in fans} and (o["fromNode"] in dnode or o["toNode"] in dnode)]
    seen = set(); k = 0
    for fb in fan_br:
        k += 1
        inlet = fb["fromNode"] if fb["toNode"] in dnode else fb["toNode"]
        start = fb["toNode"] if inlet == fb["fromNode"] else fb["fromNode"]
        chain, cur, prev, Ld = [], start, None, 0.0
        while True:  # walk the duct chain to its far end
            nxt = [A[i] for i in duct_ids if cur in (A[i]["fromNode"], A[i]["toNode"]) and A[i]["id"] not in chain]
            if not nxt: break
            b = nxt[0]; chain.append(b["id"]); Ld += b["length"]; cur = b["toNode"] if b["fromNode"] == cur else b["fromNode"]
        outlet = cur
        rest = [o for o in aw if o["id"] not in duct_ids and o["id"] != fb["id"]]
        G = collections.defaultdict(list)
        for o in rest: G[o["fromNode"]].append((o["toNode"], o["length"])); G[o["toNode"]].append((o["fromNode"], o["length"]))
        dist, back, h = {inlet: 0.0}, {}, [(0.0, inlet)]
        while h:
            dd, u = heapq.heappop(h)
            if u == outlet: break
            if dd > dist.get(u, 1e18): continue
            for v, w in G[u]:
                if dd + w < dist.get(v, 1e18): dist[v] = dd + w; back[v] = u; heapq.heappush(h, (dd + w, v))
        route = [outlet]
        while route[-1] != inlet: route.append(back[route[-1]])
        route = route[::-1]
        pts, lib, rho = curve(fans[int(fb["id"])]); pts = pts or []
        d0 = A[chain[0]]; did = f"D{k}"
        ducts_out.append({"id": did, "name": f"Duct {k} ({lib} fan, the source tool branches {', '.join(chain)})", "route": [f"SAMPLE|{n}" for n in route],
            "D_mm": round(d0["width"] * 1000), "k": d0["frictionFactor"], "kb": d0["frictionFactor"], "side": "left", "vert": "roof", "pipeLen": 3,
            "kIn": 0, "kOut": 0, "from": f"{did}_DN1", "to": f"{did}_DN2", "pointPins": {}, "legOffsets": {}, "reducers": [],
            "leak": {"on": False, "Lc": 0, "section": 20},
            "fans": [{"id": f"{did}F", "raw": 3, "fanKind": "auxiliary", "fanSource": "points",
                      "fan": {"name": lib, "source": "points", "controlMode": "curve", "points": [{"Q": q, "P": pp} for q, pp in pts], "rhoRef": rho}}]})
        notes[f"duct {did}: the source tool length {Ld:.1f} m, ASM route {dist[outlet]:.1f} m, inlet {inlet}, outlet {outlet}, no leakage in the source tool export"] += 1
        seen |= set(chain) | {fb["id"]}
    aw = [o for o in aw if o["id"] not in seen]
    fl = [f for f in fl if f["airwayId"] not in seen]
    used = {o["fromNode"] for o in aw} | {o["toNode"] for o in aw}
    out_nodes = [n for n in out_nodes if n["id"] in used]

proj = dict(projectName=pname, notes="Converted from the source tool input exports by network_export_to_asmlab.py. " + "; ".join(f"{k} ({v})" for k, v in notes.items()),
            network=dict(nodes=out_nodes, airways=aw, fans=fl, regulators=regs), nodePositions={n["id"]: dict(x=n["x"], y=n["y"]) for n in out_nodes})
os.makedirs(outdir, exist_ok=True); path = os.path.join(outdir, f"{pname}.asmlab.json")
json.dump(proj, open(path, "w"), indent=1)
if ducts_out:
    dn = [{"id": f"{d['id']}_DN1", "host": d["route"][0], "p": None} for d in ducts_out] + [{"id": f"{d['id']}_DN2", "host": d["route"][-1], "p": None} for d in ducts_out]
    json.dump({"format": "asmlab-ducts", "version": 1, "data": {"vnodes": [], "ducts": ducts_out, "ductNodes": dn, "machines": [], "assets": [],
               "seq": {"d": len(ducts_out), "c": 0, "dn": 2 * len(ducts_out), "v": 0, "m": 0, "a": 0}}}, open(path.replace(".asmlab.json", ".ducts.json"), "w"), indent=1)
    print("ducts:", path.replace(".asmlab.json", ".ducts.json"))
print(path); print(len(out_nodes), "nodes,", len(aw), "airways,", len(fl), "fans/pressure sources,", len(regs), "fixed-flow regulators")
print("boundaries:", {k: v for k, v in role.items()}); print("branch types:", dict(cnt)); print("notes:", dict(notes))
